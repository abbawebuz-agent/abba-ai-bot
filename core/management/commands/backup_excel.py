"""Barcha ma'lumotlarni Excel (.xlsx) jadval formatida → Telegram channel.

Har bir asosiy jadval alohida sheet bo'lib chiqadi (Foydalanuvchilar,
Promokodlar, Sovg'a so'rovlari, Sotuvchilar, ...). Kunlik 04:00 backup
bilan birga yuboriladi (core.tasks.daily_db_backup).

Usage:
    python manage.py backup_excel                 # excel yaratadi + yuboradi
    python manage.py backup_excel --no-upload     # faqat lokal fayl

ENV: TELEGRAM_BOT_TOKEN, BACKUP_CHANNEL_ID (backup_db bilan bir xil).
"""
import os
import tempfile
from datetime import datetime, date

from django.conf import settings
from django.core.management.base import BaseCommand, CommandError
from django.utils import timezone

from core.models import (
    TelegramUser, QRCode, QRCodeBatch, GiftRedemption, Gift,
    Seller, SellerBatch, SellerPointsTransaction, MonthlyPromoTicket,
    SellerRegistrationCode,
)
from core.management.commands.backup_db import Command as BackupCommand


# (sheet nomi, queryset) — tartib muhim emas
def _exports():
    return [
        ('Foydalanuvchilar', TelegramUser.objects.all().order_by('id')),
        ('Promokodlar', QRCode.objects.all().order_by('id')),
        ('QR partiyalari', QRCodeBatch.objects.all().order_by('id')),
        ("Sovga sorovlari", GiftRedemption.objects.all().order_by('id')),
        ("Sovgalar", Gift.objects.all().order_by('id')),
        ('Sotuvchilar', Seller.objects.all().order_by('id')),
        ('Sotuvchi partiyalari', SellerBatch.objects.all().order_by('id')),
        ('Ball tranzaksiyalari', SellerPointsTransaction.objects.all().order_by('id')),
        ('Oylik biletlar', MonthlyPromoTicket.objects.all().order_by('id')),
        ('Sotuvchi IDlari', SellerRegistrationCode.objects.all().order_by('id')),
    ]


class Command(BaseCommand):
    help = "Barcha ma'lumotlarni Excel jadval qilib Telegram kanalga yuboradi"

    def add_arguments(self, parser):
        parser.add_argument('--no-upload', action='store_true',
                            help="Telegram ga yubormaslik (faqat lokal .xlsx)")
        parser.add_argument('--out', type=str, default=None)

    def handle(self, *args, **opts):
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter

        token = getattr(settings, 'TELEGRAM_BOT_TOKEN', '') or os.environ.get('TELEGRAM_BOT_TOKEN', '')
        chat_id = (getattr(settings, 'BACKUP_CHANNEL_ID', '')
                   or os.environ.get('BACKUP_CHANNEL_ID', ''))

        if not opts['no_upload']:
            if not token:
                raise CommandError("TELEGRAM_BOT_TOKEN sozlanmagan")
            if not chat_id:
                raise CommandError("BACKUP_CHANNEL_ID sozlanmagan")

        ts = datetime.now().strftime('%Y%m%d_%H%M%S')
        out_dir = opts['out'] or tempfile.gettempdir()
        os.makedirs(out_dir, exist_ok=True)
        xlsx_path = os.path.join(out_dir, f'jip_data_{ts}.xlsx')

        wb = Workbook()
        wb.remove(wb.active)  # default bo'sh sheet o'chiriladi

        header_font = Font(bold=True, color='FFFFFF')
        header_fill = PatternFill('solid', fgColor='1F2937')
        total_rows = 0

        for sheet_name, qs in _exports():
            ws = wb.create_sheet(title=sheet_name[:31])  # Excel sheet nomi max 31
            model = qs.model
            fields = list(model._meta.concrete_fields)
            # Sarlavha
            headers = [(f.verbose_name or f.name) for f in fields]
            ws.append([str(h).capitalize() for h in headers])
            for col_i in range(1, len(headers) + 1):
                c = ws.cell(row=1, column=col_i)
                c.font = header_font
                c.fill = header_fill
                c.alignment = Alignment(vertical='center')
            # Qatorlar
            cnt = 0
            for obj in qs.iterator():
                row = [self._cell_value(obj, f) for f in fields]
                ws.append(row)
                cnt += 1
            total_rows += cnt
            # Ustun kengligi (oddiy auto)
            for col_i, f in enumerate(fields, start=1):
                width = min(40, max(12, len(str(headers[col_i - 1])) + 2))
                ws.column_dimensions[get_column_letter(col_i)].width = width
            ws.freeze_panes = 'A2'
            self.stdout.write(f"  • {sheet_name}: {cnt} qator")

        wb.save(xlsx_path)
        size_kb = os.path.getsize(xlsx_path) / 1024
        self.stdout.write(self.style.SUCCESS(
            f"✅ Excel tayyor: {os.path.basename(xlsx_path)} "
            f"({size_kb:.1f} KB, {total_rows} qator)"
        ))

        if opts['no_upload']:
            self.stdout.write(self.style.WARNING("⚠️ --no-upload — yuborilmadi"))
            return

        caption = (
            f"📊 <b>JIP Data Export (Excel)</b>\n"
            f"📅 {ts.replace('_', ' ')}\n"
            f"📋 {total_rows} qator · {size_kb:.0f} KB\n"
            f"#backup #excel #data"
        )
        bk = BackupCommand()
        bk.stdout = self.stdout
        bk._tg_upload(token, chat_id, xlsx_path, caption)
        self.stdout.write(self.style.SUCCESS("🎉 Excel kanalga yuborildi"))

    @staticmethod
    def _cell_value(obj, field):
        """FK -> str, tz-aware datetime -> naive local, bool/raw -> o'zi."""
        val = getattr(obj, field.name, None)
        if val is None:
            return ''
        if field.is_relation:
            related = getattr(obj, field.name, None)
            return str(related) if related is not None else ''
        if isinstance(val, datetime):
            if timezone.is_aware(val):
                val = timezone.localtime(val)
            return val.replace(tzinfo=None)
        if isinstance(val, date):
            return val
        if isinstance(val, (list, dict)):
            return str(val)
        return val
